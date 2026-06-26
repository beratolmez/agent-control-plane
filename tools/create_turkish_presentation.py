from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "agent-control-plane-sunum-tr.pptx"


COLORS = {
    "ink": RGBColor(25, 31, 42),
    "muted": RGBColor(91, 101, 117),
    "blue": RGBColor(48, 96, 196),
    "cyan": RGBColor(27, 154, 170),
    "green": RGBColor(52, 145, 93),
    "orange": RGBColor(221, 132, 45),
    "red": RGBColor(196, 70, 70),
    "panel": RGBColor(245, 247, 250),
    "line": RGBColor(211, 218, 228),
    "white": RGBColor(255, 255, 255),
}


def set_run(run, size=20, bold=False, color="ink"):
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = COLORS[color]


def add_title(slide, title, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.32), Inches(12.2), Inches(0.75))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.LEFT
    set_run(p.runs[0], 28, True, "ink")
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.58), Inches(0.99), Inches(11.8), Inches(0.4))
        stf = sub.text_frame
        stf.clear()
        sp = stf.paragraphs[0]
        sp.text = subtitle
        set_run(sp.runs[0], 12, False, "muted")


def add_footer(slide, n):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.24), Inches(13.33), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS["line"]
    line.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.28), Inches(12.2), Inches(0.22))
    p = box.text_frame.paragraphs[0]
    p.text = f"Agent Control Plane | Türkçe sistem sunumu | {n}"
    set_run(p.runs[0], 8, False, "muted")


def add_bullets(slide, items, x=0.7, y=1.45, w=12.0, h=5.5, size=18):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(8)
        set_run(p.runs[0], size, False, "ink")


def add_card(slide, x, y, w, h, title, body, color="blue"):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS["panel"]
    shape.line.color.rgb = COLORS["line"]
    shape.adjustments[0] = 0.08
    title_box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.14), Inches(w - 0.36), Inches(0.32))
    tp = title_box.text_frame.paragraphs[0]
    tp.text = title
    set_run(tp.runs[0], 14, True, color)
    body_box = slide.shapes.add_textbox(Inches(x + 0.18), Inches(y + 0.58), Inches(w - 0.36), Inches(h - 0.7))
    bp = body_box.text_frame.paragraphs[0]
    bp.text = body
    set_run(bp.runs[0], 12, False, "ink")


def add_flow(slide, labels, y=2.2):
    x = 0.55
    widths = [1.65, 1.8, 1.65, 1.85, 1.8, 1.65]
    colors = ["blue", "cyan", "orange", "green", "blue", "cyan"]
    for i, label in enumerate(labels):
        w = widths[i]
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.78))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLORS[colors[i]]
        shape.line.fill.background()
        shape.adjustments[0] = 0.12
        p = shape.text_frame.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        set_run(p.runs[0], 11, True, "white")
        if i < len(labels) - 1:
            arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + w + 0.08), Inches(y + 0.25), Inches(0.42), Inches(0.28))
            arr.fill.solid()
            arr.fill.fore_color.rgb = COLORS["line"]
            arr.line.fill.background()
        x += w + 0.55


def add_two_column(slide, left_title, left_items, right_title, right_items):
    add_card(slide, 0.65, 1.45, 5.9, 5.1, left_title, "\n".join(f"- {i}" for i in left_items), "blue")
    add_card(slide, 6.8, 1.45, 5.9, 5.1, right_title, "\n".join(f"- {i}" for i in right_items), "green")


slides = [
    {
        "title": "Agent Control Plane",
        "subtitle": "Pi tabanlı uzun süreli agent loop sistemi: mimari, çalışma akışı ve kullanım senaryoları",
        "type": "cover",
    },
    {
        "title": "Bu Proje Nedir?",
        "bullets": [
            "Uzun süren agent görevlerini izlenebilir döngüler halinde çalıştıran bir kontrol panelidir.",
            "Kullanıcı bir hedef verir; backend hedefi round'lara böler ve agent çalıştırır.",
            "Her run, token kullanımı, çıktı ve hata bilgisi Neon Postgres'e kaydedilir.",
            "React dashboard canlı durumu, run geçmişini ve insan onayı gereken noktaları gösterir.",
        ],
    },
    {
        "title": "Ana Bileşenler",
        "two": (
            "Backend",
            [
                "Bun + TypeScript + Hono",
                "Loop lifecycle yönetimi",
                "Pi CLI çağrısı",
                "Neon veritabanı erişimi",
            ],
            "Frontend",
            [
                "Vite + React + TypeScript",
                "Goal başlatma formu",
                "Canlı loop paneli",
                "Run history tablosu",
            ],
        ),
    },
    {
        "title": "Uçtan Uca Akış",
        "flow": [
            "Kullanıcı Goal",
            "React UI",
            "Backend API",
            "Loop Driver",
            "Pi Agent",
            "Neon + Workspace",
        ],
        "bullets": [
            "Frontend hedefi `/api/loops` endpoint'ine gönderir.",
            "Backend loop kaydını açar ve arka planda `runLoop()` çalıştırır.",
            "Pi her round için model/provider tarafında tekil agent çalıştırır.",
            "Dosyalar workspace'e; run kayıtları Neon'a yazılır.",
        ],
    },
    {
        "title": "Loop Sistemini Kim Sağlıyor?",
        "bullets": [
            "Loop sistemi Pi'nin değil, projenin backend kodunun sorumluluğundadır.",
            "`backend/src/loop.ts` gerçek döngüyü, durum geçişlerini ve worker fan-out mantığını yönetir.",
            "`backend/src/pi.ts` sadece tek bir Pi agent çağrısını başlatır ve JSON event stream'i parse eder.",
            "Pi burada model/provider/tool runtime'ı gibi davranır.",
        ],
    },
    {
        "title": "Orchestrated Mod",
        "bullets": [
            "Varsayılan moddur: önce orchestrator agent çalışır.",
            "Orchestrator read-only tool'larla mevcut ilerlemeyi inceler.",
            "Devam edilecekse 1-4 arası worker task üretir.",
            "Worker agent'lar aynı round içinde paralel çalıştırılabilir.",
            "Orchestrator `done` derse loop tamamlanır.",
        ],
    },
    {
        "title": "Ralph Mod",
        "bullets": [
            "Tek agent'lı klasik loop modelidir.",
            "Her iteration bir küçük, doğrulanabilir ilerleme yapar.",
            "Agent final mesajına `LOOP_STATUS: CONTINUE` veya `LOOP_STATUS: DONE` yazar.",
            "Backend regex ile bu sentinel satırını okuyarak devam veya bitiş kararı verir.",
        ],
    },
    {
        "title": "Model ve Provider Katmanı",
        "bullets": [
            "Backend modele doğrudan bağlanmaz; Pi CLI'a `PI_MODEL` parametresini verir.",
            "Aynı proje Groq, Mistral, OpenRouter, OpenAI/Codex, Gemini veya Ollama ile çalışabilir.",
            "Mevcut kodda orchestrator ve worker aynı `PI_MODEL` değerini kullanır.",
            "Farklı task'lara farklı model atama mevcut değildir; ama mimari olarak eklenebilir.",
        ],
    },
    {
        "title": "Senin Mevcut Kurulumun",
        "bullets": [
            "Backend portu: `http://localhost:8788`",
            "Frontend portu: `http://localhost:5173`",
            "Aktif model: `mistral/devstral-small-2507`",
            "Pi binary: `C:\\Users\\olmezpc\\AppData\\Roaming\\npm\\pi.cmd`",
            "Neon bağlantısı `.env` üzerinden sağlanıyor.",
        ],
    },
    {
        "title": "Veriler ve Dosyalar Nereye Gider?",
        "two": (
            "Diskte",
            [
                "`backend/workspaces/<loop-id>/`",
                "Agent'ın yazdığı gerçek dosyalar",
                "`PROGRESS.md` ilerleme durumu",
                "Loop'a özel çalışma alanı",
            ],
            "Neon'da",
            [
                "`loops` hedef ve durum",
                "`runs` agent çıktıları",
                "`run_events` tool/text event'leri",
                "Token, cost, hata ve session bilgisi",
            ],
        ),
    },
    {
        "title": "GitHub Repo Üzerinde Çalışma",
        "bullets": [
            "Mevcut sistem seçtiğin GitHub repo kökünde değil, izole loop workspace'inde çalışır.",
            "Repo düzenletmek için agent'a workspace içine clone yaptırabilir veya patch ürettirebilirsin.",
            "Gerçek lokal repo kökünde çalışması için projeye `target repo path` özelliği eklenmelidir.",
            "Bu özellik eklenirse dashboard'dan repo yolu seçilir ve Pi o klasörde çalıştırılır.",
        ],
    },
    {
        "title": "Retool Ne İşe Yarar?",
        "bullets": [
            "Retool zorunlu değildir; yerel React dashboard zaten çalışır.",
            "Retool, bu sistemi şirket içi/admin panel olarak yayınlamak için kullanılır.",
            "Neon'a managed Postgres resource ile bağlanabilir.",
            "Resume/stop gibi pahalı aksiyonlar permission ve confirm dialog arkasına alınabilir.",
            "Audit log, rol bazlı erişim ve operasyonel görünürlük sağlar.",
        ],
    },
    {
        "title": "API ve Kontrol Noktaları",
        "bullets": [
            "`GET /api/health`: API sağlığı ve aktif model",
            "`POST /api/loops`: yeni loop başlatma",
            "`GET /api/loops`: loop listesi",
            "`GET /api/runs`: run geçmişi",
            "`POST /api/loops/:id/pause`: current round sonrası duraklatma",
            "`POST /api/loops/:id/resume`: iteration cap sonrası devam onayı",
            "`POST /api/loops/:id/stop`: loop'u durdurma",
        ],
    },
    {
        "title": "Güvenlik ve Operasyon Notları",
        "bullets": [
            "Worker agent'lar `read,bash,edit,write` tool'larıyla gerçek dosya ve komut erişimine sahiptir.",
            "API key ve Neon connection string `.env` içinde tutulmalı, repoya commit edilmemelidir.",
            "Ücretsiz model sağlayıcılarında rate limit ve tool compatibility değişebilir.",
            "Yerel repo üzerinde çalışma açılırsa path allowlist ve sandbox politikası eklenmelidir.",
            "Port çakışmalarında backend `PORT` ve frontend `VITE_API_URL` birlikte değiştirilmelidir.",
        ],
    },
    {
        "title": "Geliştirme Fırsatları",
        "bullets": [
            "Dashboard'a hedef repo yolu seçimi eklemek.",
            "Orchestrator ve worker için ayrı model ayarı eklemek.",
            "Task büyüklüğüne göre model yönlendirme yapmak.",
            "Loop workspace diff görüntüleme ve patch export eklemek.",
            "Retool veya başka bir internal tool üzerinden rol bazlı kontrol sağlamak.",
        ],
    },
    {
        "title": "Özet",
        "bullets": [
            "Bu sistem, uzun süreli agent işlerini izlenebilir ve onaylanabilir loop'lara dönüştürür.",
            "Loop mantığı backend'dedir; Pi model/provider runtime'ıdır.",
            "Neon geçmişi, workspace ise gerçek dosyaları saklar.",
            "Groq, Mistral, OpenRouter, Gemini, OpenAI/Codex ve Ollama gibi sağlayıcılara Pi üzerinden bağlanabilir.",
            "Repo üzerinde gerçek çalışma ve gelişmiş model routing için küçük ama etkili ürünleştirme adımları gerekir.",
        ],
    },
]


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for idx, data in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank)
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLORS["white"]

        if data.get("type") == "cover":
            accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.35))
            accent.fill.solid()
            accent.fill.fore_color.rgb = COLORS["blue"]
            accent.line.fill.background()
            title = slide.shapes.add_textbox(Inches(0.75), Inches(2.15), Inches(11.7), Inches(0.8))
            p = title.text_frame.paragraphs[0]
            p.text = data["title"]
            set_run(p.runs[0], 40, True, "ink")
            subtitle = slide.shapes.add_textbox(Inches(0.78), Inches(3.05), Inches(11.4), Inches(0.8))
            sp = subtitle.text_frame.paragraphs[0]
            sp.text = data["subtitle"]
            set_run(sp.runs[0], 20, False, "muted")
            add_card(slide, 0.75, 4.35, 3.6, 1.1, "Stack", "Bun + Hono\nReact + Vite\nPi + Neon", "blue")
            add_card(slide, 4.65, 4.35, 3.6, 1.1, "Amaç", "Uzun agent işlerini\nizlenebilir hale getirmek", "green")
            add_card(slide, 8.55, 4.35, 3.6, 1.1, "Kapsam", "Kurulum, model, loop,\nveri akışı, Retool", "orange")
        else:
            add_title(slide, data["title"], data.get("subtitle"))
            if "flow" in data:
                add_flow(slide, data["flow"], y=1.55)
                add_bullets(slide, data["bullets"], y=3.05, size=17)
            elif "two" in data:
                lt, li, rt, ri = data["two"]
                add_two_column(slide, lt, li, rt, ri)
            else:
                add_bullets(slide, data["bullets"])
        add_footer(slide, idx)

    OUT.parent.mkdir(exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
